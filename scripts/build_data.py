"""Regenerates src/data/bikes.json + public/images/bikes from the Word catalogue.
Usage: python3 scripts/build_data.py BikeVault-LK-bike-list-with-photos.docx"""
import docx,re,json,sys,shutil,os
S={1:(739,23.5,'V-twin, side-valve'),2:(745,26,'Flat-twin'),3:(749,26,'Flat-twin'),4:(500,None,'V-twin, side-valve'),5:(496,13,'Single, side-valve'),6:(347,None,'Single, OHV'),7:(350,None,'Single, OHV'),8:(490,None,'Single, side-valve'),9:(746,22,'Flat-twin, side-valve'),10:(1200,None,'V-twin, side-valve'),11:(998,55,'V-twin, OHV'),12:(649,46,'Parallel-twin'),13:(828,58,'Parallel-twin'),14:(736,67,'Inline-four'),15:(499,None,'Single'),16:(987,40,'V-twin, OHV'),17:(346,None,'Single'),18:(349,20.2,'Single'),19:(199.5,24.5,'Single'),20:(155,18.4,'Single'),21:(471,47,'Parallel-twin'),22:(948,125,'Inline-four'),23:(1252,121,'V-twin'),24:(803,73,'L-twin'),25:(293,None,'Two-stroke single'),26:(155,None,'Single'),27:(None,110,'Electric motor'),28:(198,None,'Single'),29:(110,7.8,'Single'),30:(149,12.4,'Single'),31:(249,26.5,'Single'),32:(452,40,'Single'),33:(102,7.9,'Single'),34:(1000,217.5,'Inline-four'),35:(998,200,'Inline-four'),36:(998,203,'Inline-four'),37:(1103,215,'V4'),38:(999,210,'Inline-four'),39:(765,130,'Inline-triple'),40:(1301,180,'V-twin'),41:(648,47,'Parallel-twin'),42:(1084,102,'Parallel-twin'),43:(1254,136,'Boxer-twin')}
d=docx.Document(sys.argv[1]);out=[]
os.makedirs('public/images/bikes',exist_ok=True)
for r in d.tables[0].rows[1:]:
    c=r.cells;n=int(c[0].text.strip());name,brand=[p.text.strip() for p in c[1].paragraphs if p.text.strip()][0].split('\n')
    y,cat=[x.strip() for x in c[2].text.split('•')]
    rid=re.findall(r'r:embed="(rId\d+)"',c[3]._tc.xml)[0];part=d.part.related_parts[rid]
    fn=f"{n:02d}."+part.partname.split('.')[-1];open('public/images/bikes/'+fn,'wb').write(part.blob)
    typ=[p.text.strip() for p in c[3].paragraphs if p.text.strip()][0]
    cc,hp,eng=S[n];slug=re.sub(r'[^a-z0-9]+','-',f'{brand} {name}'.lower()).strip('-')
    out.append(dict(id=f'{n:02d}-{slug}',n=n,name=name,brand=brand,year=int(y),category=cat,image=fn,imageType=typ,cc=cc,powerHp=hp,engineType=eng))
json.dump(out,open('src/data/bikes.json','w'),indent=1,ensure_ascii=False);print(len(out),'bikes')
